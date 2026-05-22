"""Fill .docx templates with form data using python-docx.

Templates live in ``doc_templates/`` and contain ``{{ field_name }}``
placeholders. Placeholders may appear in body paragraphs, tables, headers,
and footers. A placeholder whose value is missing is replaced with an empty
string so the document never shows raw ``{{ ... }}`` markers.
"""

import re
from datetime import datetime
from pathlib import Path

from docx import Document

# Matches {{ field_name }} with optional surrounding whitespace.
_PLACEHOLDER_RE = re.compile(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}")


def _substitute(text, data):
    """Replace every {{ key }} in *text* with data[key] (blank if absent)."""
    return _PLACEHOLDER_RE.sub(lambda m: str(data.get(m.group(1), "")), text)


def _fill_paragraph(paragraph, data):
    """Replace placeholders in a paragraph, even when they span runs.

    python-docx splits text into runs that can break a ``{{ token }}`` across
    several runs. We join the run text, substitute, then write the result back
    into the first run and clear the others. The first run's formatting is
    preserved for the whole paragraph.
    """
    if not paragraph.runs:
        return
    full_text = "".join(run.text for run in paragraph.runs)
    if "{{" not in full_text:
        return
    new_text = _substitute(full_text, data)
    if new_text == full_text:
        return
    paragraph.runs[0].text = new_text
    for run in paragraph.runs[1:]:
        run.text = ""


def _fill_tables(tables, data):
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _fill_paragraph(paragraph, data)
                # Tables can nest.
                _fill_tables(cell.tables, data)


def _fill_headers_footers(document, data):
    for section in document.sections:
        for container in (section.header, section.footer,
                          section.first_page_header, section.first_page_footer,
                          section.even_page_header, section.even_page_footer):
            for paragraph in container.paragraphs:
                _fill_paragraph(paragraph, data)
            _fill_tables(container.tables, data)


def generate_document(template_path, data, output_path):
    """Render *template_path* with *data* and save to *output_path*.

    Returns the output ``Path``. Raises ``FileNotFoundError`` if the template
    is missing.
    """
    template_path = Path(template_path)
    if not template_path.exists():
        raise FileNotFoundError(f"Template not found: {template_path}")

    document = Document(str(template_path))

    for paragraph in document.paragraphs:
        _fill_paragraph(paragraph, data)
    _fill_tables(document.tables, data)
    _fill_headers_footers(document, data)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return output_path


def build_output_filename(matter_key, client_name):
    """Return a filesystem-safe, timestamped output filename."""
    safe_client = re.sub(r"[^A-Za-z0-9]+", "_", client_name or "client").strip("_")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"{matter_key}_{safe_client}_{timestamp}.docx"
