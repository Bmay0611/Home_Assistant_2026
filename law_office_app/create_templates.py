"""Generate the starter .docx templates in doc_templates/.

These are plain, editable Word documents containing {{ placeholder }} tokens
that match the field names in matters.py. Open them in Word to adjust wording,
formatting, captions, or letterhead -- just keep the {{ placeholders }} intact.

Run once:  python create_templates.py
Re-running overwrites the starter templates, so edit copies, not these, if you
want to keep customizations under a different name.
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

import config


def _new_doc():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    return doc


def _heading(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13)
    return p


def _caption(doc):
    """Standard court caption shared by all pleadings."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("IN THE CIRCUIT COURT OF {{ court_county }} COUNTY, "
                    "{{ court_state }}")
    run.bold = True
    doc.add_paragraph("Case No. {{ case_number }}")
    doc.add_paragraph()


def _signature_block(doc):
    doc.add_paragraph()
    doc.add_paragraph("Respectfully submitted,")
    doc.add_paragraph()
    doc.add_paragraph("_______________________________")
    doc.add_paragraph("{{ attorney_name }}")
    doc.add_paragraph("Bar No. {{ attorney_bar_number }}")
    doc.add_paragraph("{{ attorney_firm }}")
    doc.add_paragraph("{{ attorney_address }}")
    doc.add_paragraph("Phone: {{ attorney_phone }}")
    doc.add_paragraph("Email: {{ attorney_email }}")


def build_divorce_petition():
    doc = _new_doc()
    _caption(doc)
    doc.add_paragraph("{{ petitioner_name }}, Petitioner,")
    doc.add_paragraph("v.")
    doc.add_paragraph("{{ respondent_name }}, Respondent.")
    doc.add_paragraph()
    _heading(doc, "PETITION FOR DISSOLUTION OF MARRIAGE")
    doc.add_paragraph()
    doc.add_paragraph(
        "COMES NOW the Petitioner, {{ petitioner_name }}, residing at "
        "{{ petitioner_address }}, and for this Petition for Dissolution of "
        "Marriage against the Respondent, {{ respondent_name }}, residing at "
        "{{ respondent_address }}, states as follows:")
    doc.add_paragraph(
        "1. The parties were married on {{ marriage_date }} in "
        "{{ marriage_place }} and last separated on {{ separation_date }}.")
    doc.add_paragraph(
        "2. The grounds for dissolution are as follows: {{ grounds }}")
    doc.add_paragraph(
        "3. The minor children of the marriage are: {{ minor_children }}")
    doc.add_paragraph(
        "4. Petitioner requests the following relief: {{ relief_requested }}")
    doc.add_paragraph()
    doc.add_paragraph(
        "WHEREFORE, Petitioner prays that the Court dissolve the marriage and "
        "grant the relief requested above, together with such other relief as "
        "the Court deems just and proper.")
    doc.add_paragraph()
    doc.add_paragraph("Dated: {{ filing_date }}")
    _signature_block(doc)
    return doc


def build_parenting_plan():
    doc = _new_doc()
    _caption(doc)
    doc.add_paragraph("{{ parent_one_name }} and {{ parent_two_name }}")
    doc.add_paragraph()
    _heading(doc, "PARENTING PLAN")
    doc.add_paragraph()
    doc.add_paragraph(
        "This Parenting Plan is entered into between {{ parent_one_name }} "
        "(\"Parent One\") and {{ parent_two_name }} (\"Parent Two\") "
        "concerning the following child(ren):")
    doc.add_paragraph("{{ children }}")
    doc.add_paragraph()
    doc.add_paragraph("1. RESIDENTIAL SCHEDULE").runs[0].bold = True
    doc.add_paragraph("{{ residential_schedule }}")
    doc.add_paragraph("2. HOLIDAY SCHEDULE").runs[0].bold = True
    doc.add_paragraph("{{ holiday_schedule }}")
    doc.add_paragraph("3. DECISION MAKING").runs[0].bold = True
    doc.add_paragraph("{{ decision_making }}")
    doc.add_paragraph("4. TRANSPORTATION AND EXCHANGE").runs[0].bold = True
    doc.add_paragraph("{{ transportation }}")
    doc.add_paragraph()
    doc.add_paragraph("Agreed and dated: {{ agreement_date }}")
    doc.add_paragraph()
    doc.add_paragraph("_______________________________")
    doc.add_paragraph("{{ parent_one_name }}, Parent One")
    doc.add_paragraph()
    doc.add_paragraph("_______________________________")
    doc.add_paragraph("{{ parent_two_name }}, Parent Two")
    _signature_block(doc)
    return doc


def build_notice_of_hearing():
    doc = _new_doc()
    _caption(doc)
    doc.add_paragraph("{{ petitioner_name }}, Petitioner,")
    doc.add_paragraph("v.")
    doc.add_paragraph("{{ respondent_name }}, Respondent.")
    doc.add_paragraph()
    _heading(doc, "NOTICE OF HEARING")
    doc.add_paragraph()
    doc.add_paragraph(
        "PLEASE TAKE NOTICE that the above matter, {{ hearing_matter }}, will "
        "be heard before the Honorable {{ judge_name }} on {{ hearing_date }} "
        "at {{ hearing_time }}, at the following location:")
    doc.add_paragraph("{{ hearing_location }}")
    doc.add_paragraph()
    doc.add_paragraph(
        "You are advised to appear and be heard at the time and place stated "
        "above.")
    doc.add_paragraph()
    doc.add_paragraph("Dated: {{ notice_date }}")
    _signature_block(doc)
    return doc


BUILDERS = {
    "divorce_petition.docx": build_divorce_petition,
    "parenting_plan.docx": build_parenting_plan,
    "notice_of_hearing.docx": build_notice_of_hearing,
}


def main():
    config.TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    for filename, builder in BUILDERS.items():
        path = config.TEMPLATE_DIR / filename
        builder().save(str(path))
        print(f"Wrote {path}")


if __name__ == "__main__":
    main()
